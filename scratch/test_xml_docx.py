import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
table = doc.add_table(rows=2, cols=2)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcMar_xml = f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
    cell._tc.get_or_add_tcPr().append(parse_xml(tcMar_xml))

def set_cell_shading(cell, color_hex: str):
    shd_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_xml = f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="9CA3AF"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/><w:left w:val="none"/><w:right w:val="none"/><w:insideV w:val="none"/></w:tblBorders>'
    tblPr.append(parse_xml(borders_xml))

set_cell_margins(table.cell(0, 0))
set_cell_shading(table.cell(0, 0), "1F2937")
set_table_borders(table)

doc.save("scratch/test_table.docx")
print("Saved successfully!")
