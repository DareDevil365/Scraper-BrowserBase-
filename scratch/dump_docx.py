import docx
from pathlib import Path

doc_path = Path(r"c:\Users\yasha\Desktop\scout\outputs\Answer_New.docx")
print("Path exists:", doc_path.exists())
if doc_path.exists():
    doc = docx.Document(str(doc_path))
    print("Paragraph count:", len(doc.paragraphs))
    print("Table count:", len(doc.tables))
    for idx, p in enumerate(doc.paragraphs[:30]):
        print(f"P {idx}: {p.text[:100]}")
