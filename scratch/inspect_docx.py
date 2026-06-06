import docx
from pathlib import Path

run_dir = Path(r"c:\Users\yasha\Desktop\scout\outputs\create_me_a_roadmap_for_learning_coding_and_ai_ml_in_my_1_5__20260606030607_15e739fa")
docx_path = run_dir / "report.docx"
print("Docx path exists:", docx_path.exists())

if docx_path.exists():
    doc = docx.Document(str(docx_path))
    
    print("\n=== Document Sections ===")
    print("Number of sections:", len(doc.sections))
    for idx, sec in enumerate(doc.sections):
        print(f"Section {idx}: margins (top={sec.top_margin.inches}, bottom={sec.bottom_margin.inches}, left={sec.left_margin.inches}, right={sec.right_margin.inches})")

    print("\n=== Document Paragraphs (First 20) ===")
    for idx, p in enumerate(doc.paragraphs[:20]):
        runs_info = [f"'{r.text}'(bold={r.bold}, italic={r.italic}, font={r.font.name}, size={r.font.size})" for r in p.runs]
        print(f"P {idx} (style={p.style.name}): text='{p.text[:60]}' | runs={runs_info[:3]}")

    print("\n=== Document Tables ===")
    print("Number of tables:", len(doc.tables))
    for idx, t in enumerate(doc.tables):
        print(f"Table {idx}: rows={len(t.rows)}, cols={len(t.columns)}, style='{t.style.name}'")
        # Print first row cells
        row_cells = [cell.text.strip().replace('\n', ' ')[:30] for cell in t.rows[0].cells]
        print(f"  Header Row: {row_cells}")
        if len(t.rows) > 1:
            first_data = [cell.text.strip().replace('\n', ' ')[:30] for cell in t.rows[1].cells]
            print(f"  First Data Row: {first_data}")
